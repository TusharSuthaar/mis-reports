"""Writes a made-up but realistic month of data (August 2026) in the exact layouts of the real files, so the app can be
run and tested anywhere (cloud sessions included) without real client data. Nothing here comes from a real file:
first names are common ones but surnames are invented, amounts are random, and every PAN has 'X' as its 4th letter, which no real PAN
has (the 4th letter is the holder type: P, C, H, F, ...), so a demo PAN can never be a real person's.

The same seed always gives the same files, so totals can be checked exactly.

    pip install openpyxl
    python tools/demo-data/generate.py            # writes demo-data/*.xlsx and demo-data/products.json

Files (upload order is the order of a fresh set-up; tools/demo-data/seed.py does it all):
    1 masters.xlsx              Setup > Import (regions, branches, RMs, RDs)
    2 ClientMasterReport.xlsx   Data uploads > Finesse client master (Ops)
    3 Master_MIS.xlsx           Data uploads > Master MIS (partner sheets + per-client RM / partner / %)
    4 MFBrokerage_Aug2026.xlsx  Month close > Finesse MF brokerage, month 2026-08
    5 Equity_Brokerage_Aug2026.xlsx  Products > Equity upload (broker's Brokerage Due report)
    6 products.json             hand-entered FD / Bonds / AIF / Loans / GIFT / PMS / NPS / P2P entries (seed.py posts them)
"""
import json
import math
import os
import random
from datetime import date, timedelta

import openpyxl

SEED = 20260831
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', 'demo-data')
MONTH = date(2026, 8, 1)
MONTH_END = date(2026, 8, 31)
ORG_FROM = date(2025, 4, 1)

FIRST = ['Aarav', 'Aditi', 'Akash', 'Amit', 'Anjali', 'Ankit', 'Anita', 'Arjun', 'Asha', 'Deepak', 'Divya', 'Gaurav', 'Geeta', 'Harish',
         'Isha', 'Jatin', 'Kavita', 'Kiran', 'Komal', 'Lalit', 'Manish', 'Meena', 'Mohit', 'Neha', 'Nikhil', 'Nisha', 'Pankaj', 'Pooja',
         'Prakash', 'Priya', 'Rahul', 'Rajesh', 'Ravi', 'Rekha', 'Ritu', 'Rohit', 'Sachin', 'Sanjay', 'Sapna', 'Seema', 'Shalini', 'Shreya',
         'Sneha', 'Sunil', 'Sunita', 'Suresh', 'Swati', 'Tarun', 'Uma', 'Varun', 'Vidya', 'Vijay', 'Vikas', 'Vinod', 'Yash', 'Zoya']
# Invented surnames (syllable pairs), so a demo full name practically never equals a real person's.
LAST = sorted({a + b for a in ['Aru', 'Bhav', 'Chel', 'Dhar', 'Elam', 'Gauv', 'Hiral', 'Indu', 'Javr', 'Kesh', 'Lohv', 'Mirv', 'Nalv',
                               'Orv', 'Parv', 'Rudv', 'Selv', 'Tarv', 'Udav', 'Varv']
               for b in ['anthi', 'ekar', 'ivale', 'ombe', 'urkar', 'aseth', 'imani', 'oriya']})
# Firm names carry the word Demo so they can't be mistaken for a real partner firm.
FIRMS = ['%s Demo' % f for f in ['Sunrise', 'Lotus', 'Evergreen', 'Northstar', 'Pinnacle', 'Bluewater', 'Harmony', 'Silverline', 'Keystone',
                                 'Meadow', 'Summit', 'Riverbend', 'Goldleaf', 'Brightpath', 'Oakwood', 'Horizon', 'Trident', 'Banyan',
                                 'Saffron', 'Peacock', 'Monsoon', 'Indigo', 'Cedar', 'Falcon', 'Lighthouse', 'Maple', 'Orchid', 'Coral']]

# Region -> its branches. The HO SALES region holds the HEAD OFFICE house book.
REGIONS = {
    'NORTH & PUNJAB': ['DELHI', 'LUDHIANA'],
    'WEST & GUJARAT': ['AHMEDABAD', 'SURAT', 'MUMBAI'],
    'SOUTH': ['BANGALORE', 'KOCHI'],
    'EAST & BIHAR': ['PATNA', 'RANCHI'],
    'HO SALES': ['HEAD OFFICE'],
}
PRODUCTS = ['MF', 'Pre-IPO', 'FD', 'Bonds', 'Loans', 'Equity', 'PMS', 'AIF', 'Advisory', 'GIFT', 'Insurance']

rnd = random.Random(SEED)
used_pans = set()


def pan():
    """A PAN-shaped code that no real PAN can have: the 4th letter (holder type) is X."""
    while True:
        p = ''.join(rnd.choice('ABCDEFGHIJKLMNOPQRSTUVWXYZ') for _ in range(3)) + 'X' + rnd.choice('ABCDEFGHIJKLMNOPQRSTUVWXYZ') \
            + '%04d' % rnd.randint(1, 9999) + rnd.choice('ABCDEFGHIJKLMNOPQRSTUVWXYZ')
        if p not in used_pans:
            used_pans.add(p)
            return p


def person():
    return '%s %s' % (rnd.choice(FIRST), rnd.choice(LAST))


def money(lo, hi):
    return round(rnd.uniform(lo, hi), 2)


def save(wb, name):
    path = os.path.join(OUT, name)
    wb.save(path)
    print('wrote', os.path.relpath(path))


def fmt(d):
    return d.strftime('%d-%b-%Y')


# ---- org ---------------------------------------------------------------------------------------------------------------
employees = []   # dict(code, name, type, branch)
emp_no = 100
for region, branches in REGIONS.items():
    for b in branches:
        if b == 'HEAD OFFICE':
            employees.append(dict(code='HO001', name='HEAD OFFICE', type='HOUSE', branch=b, region=region))
            continue
        for _ in range(2 if b not in ('MUMBAI', 'DELHI') else 3):
            emp_no += 1
            employees.append(dict(code='EMP%d' % emp_no, name=person().upper(), type='RM', branch=b, region=region))
rds = {}
for region in REGIONS:
    rms = [e for e in employees if e['region'] == region and e['type'] == 'RM']
    rds[region] = rms[0] if rms else next(e for e in employees if e['region'] == region)
rms = [e for e in employees if e['type'] == 'RM'] + [e for e in employees if e['type'] == 'HOUSE']

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Regions'
ws.append(['Region'])
for r in REGIONS:
    ws.append([r])
ws = wb.create_sheet('Branches')
ws.append(['Branch', 'Region', 'From'])
for r, bs in REGIONS.items():
    for b in bs:
        ws.append([b, r, ORG_FROM])
ws = wb.create_sheet('Employees')
ws.append(['Emp Code', 'Name', 'Type'])
for e in employees:
    ws.append([e['code'], e['name'], e['type']])
ws = wb.create_sheet('RD')
ws.append(['Region', 'Emp Code', 'From'])
for r, e in rds.items():
    if e['type'] == 'RM':
        ws.append([r, e['code'], ORG_FROM])
ws = wb.create_sheet('RM-Branch')
ws.append(['Emp Code', 'Branch', 'From'])
for e in employees:
    ws.append([e['code'], e['branch'], ORG_FROM])

# ---- partners (Master MIS partner sheets) ---------------------------------------------------------------------------------
firms = FIRMS[:]
rnd.shuffle(firms)
sub_brokers = []
for i in range(6):
    sub_brokers.append(dict(name='%s Wealth Partners' % firms.pop(), type='Sub broker',
                            pct={'MF': rnd.choice([0.5, 0.6, 0.7, 0.8]), 'Pre-IPO': 0.5, 'FD': 0.5, 'Bonds': 0.5, 'Loans': 0.5, 'Equity': 0.4,
                                 'PMS': 0.5, 'AIF': 0.5, 'Advisory': 0.5, 'GIFT': 0.5, 'Insurance': 0}))
referrals = []
for i in range(20):
    nm = person() if i % 3 else '%s Investments' % firms.pop()
    referrals.append(dict(name=nm, type='Referral Partner',
                          pct={'MF': rnd.choice([0, 0, 0.3, 0.4]), 'Pre-IPO': 0.5, 'FD': 0.5, 'Bonds': 0.5, 'Loans': 0.5, 'Equity': 0.3,
                               'PMS': 0.5, 'AIF': 0.5, 'Advisory': 0.5, 'GIFT': 0.5, 'Insurance': 0}))
# One partner whose approval is pending: listed, linked to clients, paid nothing yet.
referrals[-1]['type'] = 'REFERRAL PARTNER (UNDER PROCESS)'

# ---- clients -------------------------------------------------------------------------------------------------------------
clients = []   # dict(no, name, pan, rm, codes, guardian_of, partners)
code_no = 400000
for i in range(2400):
    rm = rnd.choices(rms, weights=[3 if e['type'] == 'RM' else 1 for e in rms])[0]
    c = dict(no=i + 1, name=person().upper(), pan=pan(), rm=rm, codes=[], partners=[], minor=False, family_head=None)
    code_no += rnd.randint(1, 3)
    c['codes'].append('E%d' % code_no)
    if rnd.random() < 0.025:   # a second Finesse code under the same PAN (e.g. a joint or old account)
        code_no += 1
        c['codes'].append('E%d' % code_no)
    r = rnd.random()
    if r < 0.06:
        sb = rnd.choice(sub_brokers)
        c['partners'].append((sb, sb['pct']['MF']))
    elif r < 0.33:
        ref = rnd.choice(referrals)
        pct = ref['pct']['MF']
        if rnd.random() < 0.08:   # a client exception: this client's own %
            pct = rnd.choice([0.2, 0.25, 0.5])
        c['partners'].append((ref, pct))
        if rnd.random() < 0.05:   # also a sub-broker at 0% (both roles at once)
            c['partners'].insert(0, (rnd.choice(sub_brokers), 0))
    clients.append(c)
# Families: about a third of clients sit under another client as family head.
for c in clients:
    c['family_head'] = rnd.choice(clients) if rnd.random() < 0.3 else c
# Minors: no PAN of their own, mapped through the guardian's PAN (they are not rows in Master MIS).
minors = []
for g in rnd.sample(clients, 25):
    code_no += 1
    minors.append(dict(name='%s %s' % (rnd.choice(FIRST).upper(), g['name'].split()[-1]), code='E%d' % code_no, guardian=g))

# ---- Finesse client master (Ops) -----------------------------------------------------------------------------------------
HEAD = ['Client Code', 'Salutation', 'Client Name', 'Family Name', 'Family Head Code', 'Relationship with Family ', 'Family Head Name',
        'Family Head', 'RM', 'SRM', 'Pan Number', 'Email', 'Reporting Email', 'Phone Number', 'Resident Status', 'Billable Status', 'DOB',
        'AGE', 'Gender', 'Branch Name', 'Date of joining', 'Client Tag', 'Call me as', 'Employment', 'Occupation', 'Website',
        'Company Working for', 'Designation', 'Nominee Name', 'Nominee Relation', 'Nominee Contact ', 'Guardian Name', 'Guardian PAN',
        'Marital Status', 'Spouse Name', 'Anniversary Date', 'Client Type', 'Client Category', 'Individual', 'BSE Login Status',
        'BSE Mandate Status', 'NSE Login Status', 'NSE Mandate Status', 'MFU Login Status', 'MFU Mandate Status', 'Latest Agreement Date',
        'Latest Agreement Status', 'KYC Fetch Date', 'KYC Status', 'DIGIO Mandate ', 'Risk Profile', 'Risk Score', 'Latest Risk Assessment ',
        'Agreed Debt Ratio', 'Agreed Equity Ratio', 'Client Status', 'Login Status', 'Referred By', 'Holding Nature', 'Joint Holder 1 PAN',
        'Joint Holder 1 Name', 'Joint Holder 1 DOB', 'Joint Holder 2 PAN', 'Joint Holder 2 Name', 'Joint Holder 2 DOB']
CITIES = {'DELHI': 'New Delhi', 'LUDHIANA': 'Ludhiana', 'AHMEDABAD': 'Ahmedabad', 'SURAT': 'Surat', 'MUMBAI': 'Mumbai', 'BANGALORE': 'Bengaluru',
          'KOCHI': 'Kochi', 'PATNA': 'Patna', 'RANCHI': 'Ranchi', 'HEAD OFFICE': 'Mumbai'}


def master_row(code, name, pan_no, guardian_name, guardian_pan, head, rm, branch, age):
    row = dict.fromkeys(HEAD)
    born = date(2026 - age, rnd.randint(1, 12), rnd.randint(1, 28))
    row.update({'Client Code': code, 'Client Name': name, 'Family Name': head['name'].split()[-1] + ' FAMILY',
                'Family Head Code': head['codes'][0], 'Relationship with Family ': 'Self' if head['name'] == name else 'Family Member',
                'Family Head Name': head['name'], 'Family Head': 'Yes' if head['name'] == name else 'No', 'RM': rm, 'Pan Number': pan_no,
                'Email': '%s.%s@example.com' % (name.split()[0].lower(), code.lower()), 'Phone Number': '90000%05d' % rnd.randint(0, 99999),
                'Resident Status': 'Resident Individual', 'Billable Status': 'Billable', 'DOB': born.strftime('%d-%m-%Y'), 'AGE': age,
                'Branch Name': branch, 'Date of joining': fmt(date(2018 + rnd.randint(0, 7), rnd.randint(1, 12), rnd.randint(1, 28))),
                'Call me as': name.split()[0].title(), 'Guardian Name': guardian_name, 'Guardian PAN': guardian_pan,
                'Individual': 'Yes', 'BSE Login Status': 'Active', 'BSE Mandate Status': 'Approved', 'NSE Login Status': 'Active',
                'NSE Mandate Status': 'Approved', 'MFU Login Status': 'Inactive', 'MFU Mandate Status': 'Not Registered',
                'DIGIO Mandate ': 'No', 'Client Status': 'Active', 'Holding Nature': 'Single'})
    return [row[h] for h in HEAD]


wb_fin = openpyxl.Workbook()
ws = wb_fin.active
ws.title = 'ClientMasterReport'
ws.append(HEAD)
addr = wb_fin.create_sheet('Client Address Master')
addr.append(['Client Code', 'Client Name', 'Pan Number', 'Family ', 'RM', 'SRM', 'Address Tag', 'Address ', 'Address Line ', 'Address Line ',
             'Address Line ', 'Country', 'State', 'City', 'PIN Code'])
for c in clients:
    for code in c['codes']:
        ws.append(master_row(code, c['name'], c['pan'], None, None, c['family_head'], c['rm']['name'], c['rm']['branch'], rnd.randint(24, 80)))
        addr.append([code, c['name'], c['pan'], c['family_head']['name'].split()[-1] + ' FAMILY', c['rm']['name'], None, 'Residential',
                     '%d, Demo Residency' % rnd.randint(1, 400), 'Sector %d' % rnd.randint(1, 60), None, None, 'India', None,
                     CITIES[c['rm']['branch']], '4%05d' % rnd.randint(0, 99999)])
for m in minors:
    g = m['guardian']
    ws.append(master_row(m['code'], m['name'], None, g['name'], g['pan'], g, g['rm']['name'], g['rm']['branch'], rnd.randint(3, 16)))

# ---- Master MIS ----------------------------------------------------------------------------------------------------------
wb_mis = openpyxl.Workbook()
for sheet, first, plist in (('SUB_BROKER', 'Sub-Broker', sub_brokers), ('REFERRAL PARTNER', 'Referral Partner', referrals)):
    ws = wb_mis.active if sheet == 'SUB_BROKER' else wb_mis.create_sheet(sheet)
    ws.title = sheet
    ws.append([first] + PRODUCTS)
    for p in plist:
        ws.append([p['name']] + [p['pct'][k] for k in PRODUCTS])
ws = wb_mis.create_sheet('Master')
ws.append(['Client No', 'REGION', 'RD', 'BRANCH', 'RM', 'CLIENT NAME', 'PAN NO', 'Concatenate', 'CLIENT TYPE', 'BROKER', 'TYPE', '%',
           'BROKER 2', 'TYPE 2', '% 2'])
for c in clients:
    rm = c['rm']
    p = c['partners'] + [(None, None)] * (2 - len(c['partners']))
    cells = []
    for partner, pct in p:
        cells += [partner['name'] if partner else None, partner['type'] if partner else None, pct]
    ws.append([c['no'], rm['region'], rds[rm['region']]['name'], rm['branch'], rm['name'], c['name'], c['pan'], c['pan'] + c['name'],
               'OFFLINE'] + cells)

# ---- Finesse MF brokerage (the month) ------------------------------------------------------------------------------------
wb_mf = openpyxl.Workbook()
summary = wb_mf.active
summary.title = 'MFBrokerageSummary'
summary.append(['MF Brokerage Summary  01-Aug-2026 to 31-Aug-2026'])
summary.append(['AMC Wise', 'Amount', '%', 'IGST', 'CGST', 'SGST', 'Total GST'])
ws = wb_mf.create_sheet('Client Wise Summary')
ws.append(['Client Wise Summary  01-Aug-2026 to 31-Aug-2026'])
ws.append(['Client Code', 'Client Name', 'Family Head Name', 'Family Head Code', 'RM', 'SRM', 'Location/Branch', 'Brokerage Start period',
           'Brokerage End period', 'Brokerage Amount', 'Percentage', 'Referred By'])
total = 0.0
lines = []
for c in clients:
    if rnd.random() < 0.03:
        continue   # no brokerage this month
    for code in c['codes']:
        amount = round(math.exp(rnd.gauss(6.2, 1.3)), 4)   # median about 500, a long tail of large books
        if rnd.random() < 0.006:
            amount = -round(rnd.uniform(5, 300), 4)     # a reversal larger than the month's trail
        lines.append((code, c, amount))
for m in minors:
    lines.append((m['code'], dict(name=m['name'], family_head=m['guardian'], rm=m['guardian']['rm']), round(math.exp(rnd.gauss(5, 1)), 4)))
for code, c, amount in lines:
    head = c['family_head']
    # A few codes are paid in two periods of the month (one row each).
    parts = [amount] if rnd.random() > 0.02 else [round(amount * 0.4, 4), round(amount - round(amount * 0.4, 4), 4)]
    periods = [(MONTH, MONTH_END)] if len(parts) == 1 else [(MONTH, date(2026, 8, 15)), (date(2026, 8, 16), MONTH_END)]
    for part, (start, end) in zip(parts, periods):
        ws.append([code, c['name'], head['name'], head['codes'][0], c['rm']['name'], None, c['rm']['branch'], fmt(start), fmt(end), part,
                   100, None])
        total += part
ws.append(['Total', None, None, None, None, None, None, None, None, round(total, 4), None, None])
mf_rows = ws.max_row - 3

# ---- Equity: the broker's "Brokerage Due" report -------------------------------------------------------------------------
wb_eq = openpyxl.Workbook()
pivot = wb_eq.active
pivot.title = 'Sheet2'
pivot.append(['Sum of Franchise2'])
pivot.append(['PAN', None, 'Total'])
ws = wb_eq.create_sheet('Sheet1')
for title in ('DEMO SECURITIES LTD', 'Brokerage Due Report', 'Period : 01/08/2026 To 31/08/2026', None):
    ws.append([title])
ws.append(['Franchise', 'Franchise Name', 'Party', 'PAN', 'Name', 'Date', 'Exchange', 'Gross Recd', 'Level 1', 'Level 2', 'Level 3',
           'Level 4', 'Franchise', 'Area', 'Other'])
ws.append([None, None, None, None, None, None, None, 'Gross Brokerage', 'Share', 'Share', 'Share', 'Share', 'Share'])
gross_total = 0.0
traders = rnd.sample([c for c in clients if c['rm']['type'] == 'RM'], 43)
for c in traders:
    party = 'P%06d' % rnd.randint(1, 999999)
    for _ in range(rnd.randint(1, 9)):
        gross = round(rnd.uniform(40, 6000), 2)
        share = round(gross * rnd.choice([0.3, 0.4, 0.5, 0.6]), 2)
        d = MONTH + timedelta(days=rnd.randint(0, 30))
        ws.append(['DEM0001', 'EPSILON MONEY MART PRIVATE LTD .', party, c['pan'], c['name'], d.strftime('%d/%m/%Y'), rnd.choice(['NSE', 'BSE', 'NFO']),
                   gross, 0, 0, 0, 0, share, 'DEMOAREA', 'MART'])
        gross_total += gross
ws.append(['Total', None, None, None, None, None, None, round(gross_total, 2)])

# ---- hand-entered products ------------------------------------------------------------------------------------------------
with_ref = [c for c in clients if any(p['type'] == 'Referral Partner' for p, _ in c['partners'])]
plain = [c for c in clients if not c['partners']]


def pick(with_partner):
    c = rnd.choice(with_ref if with_partner else plain)
    ref = next((p for p, _ in c['partners'] if p['type'] == 'Referral Partner'), None)
    return dict(pan=c['pan'], partner=ref['name'] if with_partner and ref else None)


aug = MONTH.isoformat()
products = {'platforms': [
    {'name': 'DEMO BOND HOUSE', 'terms': {'gstInclusive': True, 'tdsPercent': 5, 'ourShare': 100, 'from': '2025-04-01'}},
    {'name': 'DEMO ALTERNATES', 'terms': {'gstInclusive': False, 'tdsPercent': 2, 'ourShare': 100, 'from': '2025-04-01'}},
], 'entries': []}
E = products['entries']
for i in range(6):
    who = pick(i % 2 == 0)
    amt = rnd.choice([100000, 250000, 500000, 1000000])
    E.append(dict(kind='fd', client=who, body=dict(month=aug, partner=who['partner'], fdName=rnd.choice(['DEMO FINANCE FD', 'SAMPLE HOUSING FD']),
                                                     fdNo='FD%05d' % rnd.randint(1, 99999), interestFrequency=rnd.choice(['CUMULATIVE', 'QUARTERLY']),
                                                     depositDate='2026-08-%02d' % rnd.randint(1, 28), periodMonths=12, maturityDate='2027-08-28',
                                                     ratePercent=rnd.choice([7.5, 8, 8.25]), amount=amt, certificate='RECEIVED', marginPercent=1)))
for i in range(4):
    who = pick(i == 0)
    E.append(dict(kind='bonds', client=who, body=dict(month=aug, partner=who['partner'], platform='DEMO BOND HOUSE',
                                                        bondName='DEMO STATE DEVELOPMENT %d' % (i + 1), isin='INE0DEMO%04d' % i, couponPercent=9.5,
                                                        payoutFrequency='QUARTERLY', investmentDate='2026-08-%02d' % (5 + i), maturityDate='2029-03-31',
                                                        quantity=1, ratePerUnit=1010000, investmentAmount=1000000, accruedInterest=10000,
                                                        markupPercent=0.4, tdsApplied=True, tdsPercent=5)))
for i in range(3):
    who = pick(True)
    E.append(dict(kind='aif', client=who, body=dict(month=aug, partner=who['partner'], sharingPercent=50 if i == 0 else None,
                                                      platform='DEMO ALTERNATES', fundName='DEMO GROWTH FUND %d' % (i + 1), commitment=10000000,
                                                      agreementDate='2026-08-%02d' % (10 + i), tenure='Five Years', markupPercent=1.3,
                                                      tdsApplied=True, tdsPercent=2, payoutTerms='Paid quarterly',
                                                      instalments=[dict(amount=2500000, paidOn='2026-08-%02d' % (20 + i))])))
for i in range(5):
    who = pick(i < 2)
    E.append(dict(kind='loans', client=who, body=dict(month=aug, partner=who['partner'], lender=rnd.choice(['DEMO BANK', 'SAMPLE NBFC']),
                                                        loanType=rnd.choice(['Home Loan', 'Personal Loan', 'LAP']),
                                                        disbursalDate='2026-08-%02d' % rnd.randint(1, 28), sanctionedAmount=rnd.choice([500000, 2500000, 4000000]),
                                                        ratePercent=9.5, payoutPercent=rnd.choice([0.5, 1, 1.5]), tdsApplied=True, tdsPercent=2,
                                                        loanClosed=i == 4)))
for i in range(3):
    who = pick(i == 0)
    E.append(dict(kind='gift', client=who, body=dict(month=aug, partner=who['partner'], fund='DEMO GIFT CITY FUND', amountUsd=rnd.choice([10000, 25000]),
                                                       fxRate=87.5, fxRateDate='2026-08-14', fxSource='Demo rate', onPortal=True, trxnStatus='DONE',
                                                       brokerage=rnd.choice([3000, 7500]), tdsApplied=True, tdsPercent=2)))
for i in range(6):
    who = pick(i < 2)
    E.append(dict(kind='pms', client=who, body=dict(month='2026-0%d-01' % (4 + i % 5), partner=who['partner'], amc='DEMO PMS AMC',
                                                      strategy='DEMO MULTICAP %d' % (i % 2 + 1), amount=rnd.choice([5000000, 7500000]),
                                                      investmentDate='2026-0%d-15' % (4 + i % 5), onPortal=True, trailPercent=1.25,
                                                      payoutFrequency='QUARTERLY')))
for i in range(4):
    E.append(dict(kind='pms-payout', client=None, body=dict(amc='DEMO PMS AMC', periodFrom='2026-04-01', periodTo='2026-06-01',
                                                              amount=money(3000, 15000), tdsApplied=True, tdsPercent=2, receivedOn='2026-08-21')))
for i in range(5):
    who = pick(False)
    E.append(dict(kind='nps', client=who, body=dict(month=aug, pop='DEMO NPS POP', investmentDate='2026-08-%02d' % (3 + i), amount=rnd.choice([100000, 400000]),
                                                      welcomeMail=True, onPortal=True, revenueRemarks='60% of POP charges', popCharges=2000,
                                                      sharePercent=60)))
for i in range(4):
    who = pick(False)
    E.append(dict(kind='p2p', client=who, body=dict(month=aug, platform='DEMO P2P', scheme='YEARLY', roiPercent=12, amount=rnd.choice([200000, 500000]),
                                                      fundsInitiatedOn='2026-08-%02d' % (6 + i), startDate='2026-09-01', quarterEndDate='2026-09-30',
                                                      onPortal=True, brokerage=rnd.choice([1500, 4000]), tdsApplied=True, tdsPercent=2)))

os.makedirs(OUT, exist_ok=True)
save(wb, 'masters.xlsx')
save(wb_fin, 'ClientMasterReport.xlsx')
save(wb_mis, 'Master_MIS.xlsx')
save(wb_mf, 'MFBrokerage_Aug2026.xlsx')
save(wb_eq, 'Equity_Brokerage_Aug2026.xlsx')
with open(os.path.join(OUT, 'products.json'), 'w') as f:
    json.dump(products, f, indent=1)
print('wrote', os.path.relpath(os.path.join(OUT, 'products.json')))
print('%d clients (%d with a partner), %d minors, %d MF rows, MF net %.2f, %d equity clients, %d product entries'
      % (len(clients), sum(1 for c in clients if c['partners']), len(minors), mf_rows, total, len(traders), len(E)))
